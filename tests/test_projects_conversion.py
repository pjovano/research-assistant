from __future__ import annotations

import contextlib
import io
import json
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from _support import LiveTree

LiveTree.activate()

import vault
from docx import Document
from docx.shared import Inches, Pt
from _writing import conversion
from _writing import pandoc as pandoc_module
from _writing.actions import DraftCut, ProjectCreation
from _writing.project_file import DRAFTS
from _writing.drafts_cli import DraftsCommandLine
from _writing.conversion import (
    DocxDiscard, DocxExport, DocxImport, ExportText, InboxDocx, StyledDocx, Trash, WordFile,
)
from _writing.docx_style import SUBMISSION, READING
from _writing.pandoc import Pandoc
from _writing.tree import ProjectTree, WritingFolders

PAPER = "Kant's Critique of Mathematical Reason"
WORD_FILE = "old draft.docx"
NOTED_PROSE = "# Opening\n\nA claim with a note.[^1]\n\n> A quoted passage.\n\n[^1]: The note's text.\n"
LINKED_PROSE = ("---\ntags:\n---\nA claim about [[Friedman - 1992 - Kant and the Exact Sciences|Friedman]] "
                "and [[Geometry]].%% a private remark %%\n\n![[figure.png]]\n\n> Quoted.[^1]\n\n[^1]: A note.\n")


class RecordingTrash(Trash):
    def __init__(self) -> None:
        self.taken: list[Path] = []

    def take(self, path: Path) -> None:
        self.taken.append(path)


class ConversionTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.root, ignore_errors=True)
        self.inbox = self.root / "inbox"
        self.exports = self.root / "exports"
        self.folders = WritingFolders(self.root / "writing")
        ProjectCreation(self.tree(), PAPER, None).apply()

    def tree(self) -> ProjectTree:
        return ProjectTree(self.folders)

    def project(self):
        return self.tree().named(PAPER)

    def word_file(self, markdown: str = NOTED_PROSE, name: str = WORD_FILE) -> Path:
        self.inbox.mkdir(exist_ok=True)
        path = self.inbox / name
        subprocess.run(["pandoc", "--from", "commonmark_x", "--to", "docx", "--output", str(path)],
                       input=markdown.encode("utf-8"), check=True)
        return path

    def imported(self, section: str | None = None) -> DocxImport:
        return DocxImport(self.project(), section, WordFile(self.word_file(), Pandoc()))

    def draft(self, text: str) -> Path:
        DraftCut(self.project(), None).apply()
        path = self.folders.drafts / f"{PAPER} - v01.md"
        path.write_text(text, encoding="utf-8")
        return path

    def exported(self, style=SUBMISSION, text: str = LINKED_PROSE) -> Path:
        DocxExport(self.draft(text), self.exports, StyledDocx(style, Pandoc())).apply()
        return self.exports / f"{PAPER} - v01.docx"

    def read_back(self, docx: Path) -> str:
        return Pandoc().markdown_of(docx)


class TestDocxImport(ConversionTestCase):
    def test_the_word_file_becomes_the_first_version_of_the_draft(self) -> None:
        self.imported().apply()
        self.assertTrue((self.folders.drafts / f"{PAPER} - v01.md").is_file())

    def test_a_word_footnote_becomes_a_markdown_footnote(self) -> None:
        self.imported().apply()
        text = (self.folders.drafts / f"{PAPER} - v01.md").read_text(encoding="utf-8")
        self.assertIn("A claim with a note.[^1]", text)
        self.assertIn("[^1]: The note's text.", text.replace("’", "'"))

    def test_a_heading_and_a_block_quotation_keep_their_markdown_form(self) -> None:
        self.imported().apply()
        text = (self.folders.drafts / f"{PAPER} - v01.md").read_text(encoding="utf-8")
        self.assertIn("# Opening\n", text)
        self.assertIn("> A quoted passage.\n", text)

    def test_the_draft_carries_line_feeds_alone(self) -> None:
        self.imported().apply()
        self.assertNotIn(b"\r", (self.folders.drafts / f"{PAPER} - v01.md").read_bytes())

    def test_an_import_onto_a_drafted_project_takes_the_next_version(self) -> None:
        DraftCut(self.project(), None).apply()
        self.imported().apply()
        self.assertTrue((self.folders.drafts / f"{PAPER} - v02.md").read_text(encoding="utf-8"))

    def test_a_section_import_is_versioned_apart_from_the_whole_draft(self) -> None:
        self.imported("Introduction").apply()
        self.assertTrue((self.folders.drafts / f"{PAPER} - Introduction - v01.md").is_file())

    def test_the_imported_draft_is_listed_first_under_drafts(self) -> None:
        self.imported().apply()
        self.assertEqual(self.project().read().linked(DRAFTS), [f"{PAPER} - v01"])

    def test_the_plan_counts_the_footnotes_and_writes_nothing(self) -> None:
        plan = self.imported().plan()
        self.assertEqual((plan["footnotes"], plan["images_not_carried"]), (1, 0))
        self.assertFalse(self.folders.drafts.exists())

    def test_the_import_leaves_the_word_file_in_the_inbox(self) -> None:
        self.imported().apply()
        self.assertTrue((self.inbox / WORD_FILE).is_file())


class TestInboxDocx(ConversionTestCase):
    def test_a_word_file_of_the_inbox_is_found_by_its_name(self) -> None:
        self.assertEqual(InboxDocx(self.inbox, self.word_file().name).checked(), self.inbox / WORD_FILE)

    def test_a_path_outside_the_inbox_is_read_as_a_name_in_the_inbox(self) -> None:
        outside = self.root / "elsewhere.docx"
        outside.write_bytes(b"")
        with self.assertRaisesRegex(ValueError, "no file named elsewhere.docx"):
            InboxDocx(self.inbox, str(outside)).checked()

    def test_a_file_of_another_kind_is_refused(self) -> None:
        with self.assertRaisesRegex(ValueError, "is not a .docx file"):
            InboxDocx(self.inbox, "paper.pdf").checked()


class TestDocxDiscard(ConversionTestCase):
    def test_the_word_file_goes_to_the_trash(self) -> None:
        trash = RecordingTrash()
        DocxDiscard(self.word_file(), trash).apply()
        self.assertEqual(trash.taken, [self.inbox / WORD_FILE])

    def test_the_plan_names_the_file_and_moves_nothing(self) -> None:
        trash = RecordingTrash()
        plan = DocxDiscard(self.word_file(), trash).plan()
        self.assertEqual((plan, trash.taken), ({"moved_to_trash": str(self.inbox / WORD_FILE)}, []))

    def test_a_machine_without_a_trash_command_leaves_the_file(self) -> None:
        with mock.patch.object(conversion.shutil, "which", return_value=None):
            with self.assertRaisesRegex(ValueError, "stays where it is"):
                Trash().take(self.word_file())
        self.assertTrue((self.inbox / WORD_FILE).is_file())


class TestExportText(unittest.TestCase):
    def test_the_frontmatter_is_left_out(self) -> None:
        self.assertNotIn("tags", ExportText(LINKED_PROSE).markdown())

    def test_a_wikilink_becomes_the_words_it_shows(self) -> None:
        self.assertIn("A claim about Friedman and Geometry.", ExportText(LINKED_PROSE).markdown())

    def test_an_obsidian_comment_and_an_embed_are_left_out(self) -> None:
        markdown = ExportText(LINKED_PROSE).markdown()
        self.assertNotIn("private remark", markdown)
        self.assertNotIn("figure.png", markdown)


class TestDocxExport(ConversionTestCase):
    def test_the_docx_takes_the_notes_basename(self) -> None:
        self.assertTrue(self.exported().is_file())

    def test_a_markdown_footnote_becomes_a_word_footnote(self) -> None:
        self.assertIn("[^1]: A note.", self.read_back(self.exported()))

    def test_the_export_leaves_the_draft_as_it_was(self) -> None:
        self.exported()
        self.assertEqual((self.folders.drafts / f"{PAPER} - v01.md").read_text(encoding="utf-8"), LINKED_PROSE)

    def test_the_submission_style_sets_its_font_spacing_and_margins(self) -> None:
        document = Document(str(self.exported(SUBMISSION)))
        normal, body = document.styles["Normal"], document.styles["Body Text"]
        self.assertEqual((normal.font.name, normal.font.size, body.paragraph_format.line_spacing),
                         ("Times New Roman", Pt(12), 2.0))
        self.assertEqual(document.sections[0].left_margin, Inches(1.0))

    def test_the_reading_style_differs_from_the_submission_style(self) -> None:
        document = Document(str(self.exported(READING)))
        self.assertEqual((document.styles["Normal"].font.name, document.styles["Normal"].font.size),
                         (READING.font, Pt(READING.body_points)))

    def test_no_style_keeps_a_theme_font(self) -> None:
        styles = Document(str(self.exported())).styles.element
        themed = [fonts for fonts in styles.iter() if any("Theme" in name or "cstheme" in name for name in fonts.attrib)
                  and fonts.tag.endswith("rFonts")]
        self.assertEqual(themed, [])

    def test_every_page_carries_its_number(self) -> None:
        footer = Document(str(self.exported())).sections[0].footer
        self.assertIn('w:instr="PAGE"', footer.paragraphs[0]._p.xml)

    def test_an_existing_docx_is_never_overwritten(self) -> None:
        self.exported()
        with self.assertRaisesRegex(ValueError, "name another directory"):
            DocxExport(self.folders.drafts / f"{PAPER} - v01.md", self.exports, StyledDocx(SUBMISSION, Pandoc()))

    def test_the_plan_names_the_docx_and_writes_nothing(self) -> None:
        plan = DocxExport(self.draft(LINKED_PROSE), self.exports, StyledDocx(READING, Pandoc())).plan()
        self.assertEqual((plan["style"], self.exports.exists()), ("reading", False))


class TestPandoc(unittest.TestCase):
    def test_a_machine_without_pandoc_is_told_how_to_install_it(self) -> None:
        with mock.patch.object(pandoc_module.shutil, "which", return_value=None):
            with self.assertRaisesRegex(ValueError, "brew install pandoc"):
                Pandoc().default_reference()

    def test_a_file_pandoc_cannot_read_is_refused_with_its_reason(self) -> None:
        with tempfile.TemporaryDirectory() as scratch:
            broken = Path(scratch) / "broken.docx"
            broken.write_bytes(b"not a word file")
            with self.assertRaisesRegex(ValueError, "pandoc failed"):
                Pandoc().markdown_of(broken)


class TestConversionCommands(ConversionTestCase):
    def run_command(self, *argv: str) -> object:
        out = io.StringIO()
        with mock.patch.object(vault, "WRITING_DIR", self.folders.projects.parent), \
                mock.patch.object(vault, "INBOX_DIR", self.inbox), contextlib.redirect_stdout(out):
            DraftsCommandLine().run(list(argv))
        return json.loads(out.getvalue())

    def test_import_without_apply_writes_nothing(self) -> None:
        self.word_file()
        self.assertFalse(self.run_command("import", WORD_FILE, PAPER)["applied"])
        self.assertFalse(self.folders.drafts.exists())

    def test_import_with_apply_writes_the_draft(self) -> None:
        self.word_file()
        self.run_command("import", WORD_FILE, PAPER, "--apply")
        self.assertTrue((self.folders.drafts / f"{PAPER} - v01.md").is_file())

    def test_discard_without_apply_leaves_the_word_file(self) -> None:
        self.word_file()
        self.run_command("discard", WORD_FILE)
        self.assertTrue((self.inbox / WORD_FILE).is_file())

    def test_export_takes_the_submission_style_unless_told(self) -> None:
        self.draft(LINKED_PROSE)
        plan = self.run_command("export", f"{PAPER} - v01", "--to", str(self.exports))
        self.assertEqual(plan["style"], "submission")

    def test_export_of_an_unknown_note_exits_with_its_reason(self) -> None:
        with self.assertRaisesRegex(SystemExit, "no draft, published piece, or note named"):
            self.run_command("export", "Nowhere", "--to", str(self.exports))


if __name__ == "__main__":
    unittest.main()
